"""
Packaging Sales & Inventory Automation
Main libraries: pandas and openpyxl
Run: python automation.py
Input: source_sales_data.xlsx (must be in the same folder)
Output: reports/packaging_sales_report.xlsx
"""

from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
INPUT_FILE = BASE_DIR / "source_sales_data.xlsx"
OUTPUT_DIR = BASE_DIR / "reports"
OUTPUT_FILE = OUTPUT_DIR / "packaging_sales_report.xlsx"


def load_source_data():
    """Load all source worksheets into pandas DataFrames."""
    if not INPUT_FILE.exists():
        raise FileNotFoundError(
            f"Could not find {INPUT_FILE.name}. Keep it in the same folder as automation.py."
        )
    orders = pd.read_excel(INPUT_FILE, sheet_name="Orders")
    products = pd.read_excel(INPUT_FILE, sheet_name="Products")
    customers = pd.read_excel(INPUT_FILE, sheet_name="Customers")
    return orders, products, customers


def clean_orders(orders):
    """Standardize text, remove duplicates, and flag invalid rows."""
    orders = orders.copy()
    orders["Order_Date"] = pd.to_datetime(orders["Order_Date"], errors="coerce")
    text_columns = ["Order_ID", "Customer_ID", "Customer_Name", "City",
                    "Product_ID", "Product_Name", "Payment_Status", "Notes"]
    for column in text_columns:
        orders[column] = orders[column].fillna("").astype(str).str.strip()

    # Consistent title/case formatting for selected fields
    orders["Customer_Name"] = orders["Customer_Name"].str.replace(r"\s+", " ", regex=True).str.title()
    orders["City"] = orders["City"].str.replace(r"\s+", " ", regex=True).str.title()
    orders["Payment_Status"] = orders["Payment_Status"].str.title()

    orders["Quantity"] = pd.to_numeric(orders["Quantity"], errors="coerce")
    orders["Unit_Price"] = pd.to_numeric(orders["Unit_Price"], errors="coerce")
    orders["Discount_Rate"] = pd.to_numeric(orders["Discount_Rate"], errors="coerce").fillna(0)
    orders["Amount_Received"] = pd.to_numeric(orders["Amount_Received"], errors="coerce").fillna(0)

    orders = orders.drop_duplicates(subset=["Order_ID"], keep="first")
    orders["Data_Issue"] = ""
    orders.loc[orders["Order_Date"].isna(), "Data_Issue"] += "Missing/invalid date; "
    orders.loc[orders["Quantity"].isna() | (orders["Quantity"] <= 0), "Data_Issue"] += "Invalid quantity; "
    orders.loc[orders["Unit_Price"].isna() | (orders["Unit_Price"] <= 0), "Data_Issue"] += "Invalid unit price; "
    orders.loc[~orders["Payment_Status"].isin(["Paid", "Pending", "Overdue"]), "Data_Issue"] += "Unknown payment status; "

    # Calculate values only when inputs are valid. Invalid rows remain visible for review.
    valid = (
        orders["Order_Date"].notna()
        & orders["Quantity"].notna() & (orders["Quantity"] > 0)
        & orders["Unit_Price"].notna() & (orders["Unit_Price"] > 0)
    )
    orders["Gross_Sales"] = (orders["Quantity"] * orders["Unit_Price"]).round(2)
    orders["Discount_Amount"] = (orders["Gross_Sales"] * orders["Discount_Rate"]).round(2)
    orders["Net_Sales"] = (orders["Gross_Sales"] - orders["Discount_Amount"]).round(2)
    orders["Amount_Due"] = (orders["Net_Sales"] - orders["Amount_Received"]).clip(lower=0).round(2)
    orders.loc[~valid, ["Gross_Sales", "Discount_Amount", "Net_Sales", "Amount_Due"]] = pd.NA
    orders["Valid_Record"] = valid.map({True: "Yes", False: "Review"})
    return orders


def build_reports(cleaned, products, customers):
    """Create management summaries using pandas."""
    valid = cleaned[cleaned["Valid_Record"] == "Yes"].copy()
    valid["Gross_Profit_Estimate"] = (
        valid["Net_Sales"] - valid["Quantity"] * valid["Product_ID"].map(
            products.set_index("Product_ID")["Unit_Cost"]
        )
    ).round(2)

    monthly = valid.assign(Month=valid["Order_Date"].dt.to_period("M").astype(str)).groupby(
        "Month", as_index=False
    ).agg(
        Orders=("Order_ID", "nunique"),
        Units_Sold=("Quantity", "sum"),
        Gross_Sales=("Gross_Sales", "sum"),
        Discounts=("Discount_Amount", "sum"),
        Net_Sales=("Net_Sales", "sum"),
        Amount_Received=("Amount_Received", "sum"),
        Amount_Due=("Amount_Due", "sum"),
    )

    product_summary = valid.groupby(["Product_ID", "Product_Name"], as_index=False).agg(
        Orders=("Order_ID", "nunique"),
        Units_Sold=("Quantity", "sum"),
        Gross_Sales=("Gross_Sales", "sum"),
        Discounts=("Discount_Amount", "sum"),
        Net_Sales=("Net_Sales", "sum"),
        Amount_Due=("Amount_Due", "sum"),
    )
    cost_lookup = products.set_index("Product_ID")["Unit_Cost"]
    product_summary["Estimated_Cost"] = product_summary["Units_Sold"] * product_summary["Product_ID"].map(cost_lookup)
    product_summary["Estimated_Gross_Profit"] = (product_summary["Net_Sales"] - product_summary["Estimated_Cost"]).round(2)
    product_summary = product_summary.sort_values("Net_Sales", ascending=False)

    customer_summary = valid.groupby(["Customer_ID", "Customer_Name", "City"], as_index=False).agg(
        Orders=("Order_ID", "nunique"),
        Units_Bought=("Quantity", "sum"),
        Net_Sales=("Net_Sales", "sum"),
        Amount_Received=("Amount_Received", "sum"),
        Amount_Due=("Amount_Due", "sum"),
    ).sort_values("Net_Sales", ascending=False)

    stock = products.copy()
    stock["Stock_Status"] = stock.apply(
        lambda row: "REORDER" if row["Stock_Units"] <= row["Reorder_Level"] else "OK", axis=1
    )
    low_stock = stock[stock["Stock_Status"] == "REORDER"].sort_values("Stock_Units")

    payment = valid.groupby("Payment_Status", as_index=False).agg(
        Orders=("Order_ID", "nunique"),
        Net_Sales=("Net_Sales", "sum"),
        Amount_Received=("Amount_Received", "sum"),
        Amount_Due=("Amount_Due", "sum"),
    )

    kpis = pd.DataFrame([
        {"Metric": "Valid Orders", "Value": int((cleaned["Valid_Record"] == "Yes").sum())},
        {"Metric": "Orders Requiring Review", "Value": int((cleaned["Valid_Record"] == "Review").sum())},
        {"Metric": "Net Sales", "Value": round(valid["Net_Sales"].sum(), 2)},
        {"Metric": "Estimated Gross Profit", "Value": round(valid["Gross_Profit_Estimate"].sum(), 2)},
        {"Metric": "Outstanding Receivables", "Value": round(valid["Amount_Due"].sum(), 2)},
        {"Metric": "Units Sold", "Value": int(valid["Quantity"].sum())},
        {"Metric": "Products Requiring Reorder", "Value": int(len(low_stock))},
    ])
    return {
        "Executive Summary": kpis,
        "Monthly Sales": monthly,
        "Product Performance": product_summary,
        "Customer Performance": customer_summary,
        "Payment Analysis": payment,
        "Low Stock Alerts": low_stock,
        "Cleaned Orders": cleaned,
    }


def format_workbook(path):
    """Apply professional formatting using openpyxl."""
    wb = load_workbook(path)
    navy = "17365D"
    teal = "0F6B78"
    light_blue = "D9EAF7"
    pale_yellow = "FFF2CC"
    pale_red = "FCE4D6"
    white = "FFFFFF"
    thin_grey = Side(style="thin", color="D9E1F2")

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 30
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.font = Font(name="Aptos", bold=True, color=white, size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=Side(style="medium", color=teal))

        for col_cells in ws.columns:
            max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells[:250])
            width = min(max(max_len + 2, 12), 34)
            ws.column_dimensions[get_column_letter(col_cells[0].column)].width = width

        # Shade rows requiring attention
        headers = {cell.value: cell.column for cell in ws[1]}
        if "Valid_Record" in headers:
            valid_col = headers["Valid_Record"]
            for row in range(2, ws.max_row + 1):
                if ws.cell(row, valid_col).value == "Review":
                    for cell in ws[row]:
                        cell.fill = PatternFill("solid", fgColor=pale_red)
        if "Stock_Status" in headers:
            status_col = headers["Stock_Status"]
            for row in range(2, ws.max_row + 1):
                if ws.cell(row, status_col).value == "REORDER":
                    for cell in ws[row]:
                        cell.fill = PatternFill("solid", fgColor=pale_yellow)

        # Currency-like columns
        currency_words = ("Sales", "Price", "Cost", "Profit", "Discount", "Received", "Due", "Value")
        for col_idx in range(1, ws.max_column + 1):
            header = str(ws.cell(1, col_idx).value or "")
            if any(word in header for word in currency_words) and header not in ("Units_Sold", "Units_Bought"):
                for row_idx in range(2, ws.max_row + 1):
                    ws.cell(row_idx, col_idx).number_format = '₹#,##0.00;[Red]-₹#,##0.00'
        if "Order_Date" in headers:
            col = headers["Order_Date"]
            for row_idx in range(2, ws.max_row + 1):
                ws.cell(row_idx, col).number_format = "dd-mmm-yyyy"

    # Add a useful note to the executive summary
    ws = wb["Executive Summary"]
    note_row = ws.max_row + 2
    ws.cell(note_row, 1, "Notes")
    ws.cell(note_row, 1).font = Font(bold=True, color=white)
    ws.cell(note_row, 1).fill = PatternFill("solid", fgColor=teal)
    ws.cell(note_row, 2, "Synthetic demo data. Estimated profit uses listed unit costs and excludes overheads/tax.")
    ws.cell(note_row, 2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[note_row].height = 34
    ws.column_dimensions["B"].width = max(ws.column_dimensions["B"].width, 68)
    wb.save(path)


def main():
    OUTPUT_DIR.mkdir(exist_ok=True)
    orders, products, customers = load_source_data()
    cleaned = clean_orders(orders)
    reports = build_reports(cleaned, products, customers)

    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        for sheet_name, dataframe in reports.items():
            dataframe.to_excel(writer, sheet_name=sheet_name, index=False)

    format_workbook(OUTPUT_FILE)
    print("Automation completed successfully!")
    print(f"Input file : {INPUT_FILE}")
    print(f"Output file: {OUTPUT_FILE}")
    print(f"Valid orders: {(cleaned['Valid_Record'] == 'Yes').sum()}")
    print(f"Rows to review: {(cleaned['Valid_Record'] == 'Review').sum()}")
    print("Open the report workbook and review the Executive Summary and Low Stock Alerts sheets.")


if __name__ == "__main__":
    main()
